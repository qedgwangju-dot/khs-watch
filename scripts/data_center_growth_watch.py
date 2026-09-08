#!/usr/bin/env python3

import io
import json
from pathlib import Path

import requests
from openpyxl import load_workbook

URL = "https://www.census.gov/construction/c30/xlsx/privsatime.xlsx"
OUT = Path("out")
OUT.mkdir(exist_ok=True)

r = requests.get(URL, headers={"User-Agent": "khs-watch/1.0"}, timeout=40)
r.raise_for_status()
wb = load_workbook(io.BytesIO(r.content), data_only=True)

matches = []
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for cell in row:
            text = str(cell.value or "").strip().lower()
            if "data center" in text or "datacenter" in text:
                r0 = max(1, cell.row - 4)
                r1 = min(ws.max_row, cell.row + 4)
                c0 = max(1, cell.column - 4)
                c1 = min(ws.max_column, cell.column + 18)
                block = []
                for rr in range(r0, r1 + 1):
                    block.append([ws.cell(rr, cc).value for cc in range(c0, c1 + 1)])
                matches.append({
                    "sheet": ws.title,
                    "cell": cell.coordinate,
                    "label": cell.value,
                    "bounds": [r0, r1, c0, c1],
                    "block": block,
                })

payload = {"sheets": wb.sheetnames, "matches": matches}
(OUT / "data_center_growth_diagnostic.json").write_text(
    json.dumps(payload, ensure_ascii=False, default=str, indent=2), encoding="utf-8"
)
print(json.dumps(payload, ensure_ascii=False, default=str, indent=2))
if not matches:
    raise SystemExit("Data center label not found")
