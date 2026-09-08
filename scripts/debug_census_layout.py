import io, requests
from openpyxl import load_workbook
u='https://www.census.gov/construction/c30/xlsx/privsatime.xlsx'
r=requests.get(u,timeout=40); r.raise_for_status()
w=load_workbook(io.BytesIO(r.content),data_only=True)['Private SA']
for row in range(1,23):
    print(row,[w.cell(row,c).value for c in range(1,15)])
